"""文档解析:各种格式 → 纯文本 或 结构化行。

**为什么不做成「全都转成一段文本」**:商品参数表最有效的索引形态是
「一行商品 → 一句自然语言」(见 `indexing/chunking.py` 的 `rows_to_sentences`),
而不是把整张表拍平成一坨。所以解析层要把**表格类**和**叙述类**分开对待,
产出的类型就分成 `rows` 与 `text` 两支,由摄取服务分流。

解析器的容错原则:**能读多少读多少,读不了要说清楚为什么**。
一份 30 页的 PDF 里有一页是乱码,不该让整篇文档失败——用户要的是能检索,
不是一份完美的副本。
"""

from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.core.exceptions import IngestError
from app.core.logging import log_extra

logger = logging.getLogger(__name__)

# 支持的扩展名 → 分类。分类决定走哪条解析分支、以及产出是 text 还是 rows。
TABULAR_EXTENSIONS = {".xlsx", ".xlsm", ".csv"}
TEXT_EXTENSIONS = {".md", ".markdown", ".txt", ".html", ".htm", ".pdf", ".docx"}
# 明确不支持但用户很可能上传的格式。单独列出来是为了给出**可操作**的错误提示
# ——「不支持 .doc」远不如「请在 Word 里另存为 .docx」有用。
LEGACY_UNSUPPORTED = {
    ".doc": "旧版 .doc 是二进制格式,python-docx 读不了。请在 Word 里「另存为」.docx 后再上传。",
    ".xls": "旧版 .xls 是二进制格式,openpyxl 读不了。请用 Excel 另存为 .xlsx 后再上传。",
    ".ppt": "暂不支持 PPT。请先导出为 PDF 或 Markdown。",
    ".pptx": "暂不支持 PPT。请先导出为 PDF 或 Markdown。",
}

SUPPORTED_EXTENSIONS = TABULAR_EXTENSIONS | TEXT_EXTENSIONS

# CSV / TXT 的候选编码,按命中概率排序。
#
# **这张表是中文电商场景的刚需,不是过度设计。** 国内用户导出 CSV 的默认
# 动作是「Excel → 另存为 CSV」,而中文版 Excel 写出的是 **GBK**,
# 不是 UTF-8。按 UTF-8 读会直接 UnicodeDecodeError 或读出一片乱码,
# 而"乱码"这种失败还特别隐蔽——不报错,只是检索不到东西。
#
# gb18030 是 GBK/GB2312 的超集,放一个就够覆盖国标系列。
ENCODING_CANDIDATES = ("utf-8-sig", "utf-8", "gb18030", "big5", "latin-1")


@dataclass(slots=True)
class ParsedDocument:
    """解析结果。`rows` 与 `text` 至少有一个非空。

    `warnings` 收集「没到失败、但用户应该知道」的情况,例如 PDF 里某一页
    解析不出文字(多半是扫描图)。这些会随文档状态一起回给前端。
    """

    text: str = ""
    rows: list[dict[str, Any]] = field(default_factory=list)
    meta: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    @property
    def is_tabular(self) -> bool:
        return bool(self.rows)


def classify(filename: str) -> str:
    """返回 `tabular` / `text` / 抛 `IngestError`。"""
    ext = Path(filename).suffix.lower()
    if ext in TABULAR_EXTENSIONS:
        return "tabular"
    if ext in TEXT_EXTENSIONS:
        return "text"
    if ext in LEGACY_UNSUPPORTED:
        raise IngestError(LEGACY_UNSUPPORTED[ext], filename=filename, extension=ext)
    supported = "、".join(sorted(SUPPORTED_EXTENSIONS))
    raise IngestError(
        f"不支持的文件类型 {ext or '(无扩展名)'}。当前支持:{supported}",
        filename=filename,
        extension=ext,
    )


def _decode(raw: bytes, filename: str) -> tuple[str, str]:
    """按候选编码依次尝试。返回 (文本, 实际使用的编码)。"""
    for enc in ENCODING_CANDIDATES:
        try:
            # latin-1 能解码**任何**字节序列,放在最后做兜底。它不会报错,
            # 但也几乎一定是错的——所以走到它时要留个警告。
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    # 理论上到不了这里(latin-1 永不失败),留作防御。
    raise IngestError("文件编码无法识别", filename=filename)  # pragma: no cover


def _read_text_file(path: Path, filename: str) -> ParsedDocument:
    text, encoding = _decode(path.read_bytes(), filename)
    warnings: list[str] = []
    if encoding == "latin-1":
        warnings.append(
            "文件编码未能识别,已按 latin-1 强行读取,内容可能是乱码。"
            "建议用 UTF-8 重新保存。"
        )
    return ParsedDocument(text=text, meta={"encoding": encoding}, warnings=warnings)


# ---------------------------------------------------------------------------
# 表格类
# ---------------------------------------------------------------------------


def _rows_from_matrix(matrix: list[list[Any]], *, sheet: str) -> list[dict[str, Any]]:
    """二维表 → list[dict]。第一非空行为表头。"""
    header: list[str] | None = None
    rows: list[dict[str, Any]] = []

    for raw_row in matrix:
        cells = ["" if c is None else str(c).strip() for c in raw_row]
        if not any(cells):
            continue  # 整行皆空:Excel 里大量这种行,不跳过会污染结果
        if header is None:
            # 表头里的空列用 col{i} 占位,否则 dict 会丢失这些列
            header = [c or f"col{i}" for i, c in enumerate(cells)]
            continue
        rows.append({header[i]: cells[i] for i in range(min(len(header), len(cells)))})

    for row in rows:
        row["_sheet"] = sheet
    return rows


def _count_formulas(path: Path) -> int:
    """数一下工作簿里有多少公式单元格。

    单独读一遍是为了**发现一个否则会静默发生的错误**:`data_only=True`
    读到的是公式的**缓存计算结果**,而缓存只在 Excel/LibreOffice 保存过
    该文件时才存在。程序生成的 xlsx(或从其他系统导出的)通常没有缓存,
    此时公式单元格会读成 **None**——价格列整列变空,而文档状态却是
    「成功」。这种失败最坏:用户以为入库了,问答时模型只是不提价格。
    """
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=False)
    except Exception:  # noqa: BLE001 - 探测失败不该阻断主流程
        return 0
    try:
        count = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                count += sum(
                    1 for c in row if isinstance(c, str) and c.startswith("=")
                )
                if count > 50:  # 只需知道「有没有」,不必数完
                    return count
        return count
    finally:
        wb.close()


def _parse_xlsx(path: Path, filename: str) -> ParsedDocument:
    from openpyxl import load_workbook

    warnings: list[str] = []
    formula_count = _count_formulas(path)
    if formula_count:
        warnings.append(
            f"工作簿含 {formula_count} 个公式单元格。已按公式的**计算结果**读取;"
            f"若这些单元格读出来是空的,说明该文件从未被 Excel 保存过、没有缓存值,"
            f"请用 Excel 打开后另存为一次再上传。"
        )

    try:
        # read_only=True:流式读取,几万行的商品表不会把内存吃光(本机只有 15GB)。
        # data_only=True:取公式的**计算结果**而不是公式本身——
        # 售价列如果是 =B2*1.13,不设这个会读到 "=B2*1.13" 这个字符串,
        # 然后模型会一本正经地告诉顾客价格是「=B2*1.13」。
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl 抛的异常类型很杂
        raise IngestError(f"Excel 解析失败:{exc}", filename=filename) from exc

    all_rows: list[dict[str, Any]] = []
    sheet_names: list[str] = []
    try:
        for ws in wb.worksheets:
            sheet_names.append(ws.title)
            all_rows.extend(_rows_from_matrix([list(r) for r in ws.iter_rows(values_only=True)],
                                              sheet=ws.title))
    finally:
        # read_only 模式下不 close 会留下文件句柄,Windows 上表现为后续
        # 删除该文件时报 WinError 32。
        wb.close()

    return ParsedDocument(
        rows=all_rows,
        meta={"sheets": sheet_names, "row_count": len(all_rows)},
        warnings=warnings,
    )


def _parse_csv(path: Path, filename: str) -> ParsedDocument:
    text, encoding = _decode(path.read_bytes(), filename)
    warnings: list[str] = []
    if encoding == "latin-1":
        warnings.append("CSV 编码未能识别,已按 latin-1 强行读取,内容可能是乱码。")

    # 分隔符嗅探:中文 Excel 导出的「CSV」有时是分号或制表符分隔。
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel  # 嗅探失败就按标准逗号,不让它中断解析

    matrix = [row for row in csv.reader(io.StringIO(text), dialect)]
    rows = _rows_from_matrix(matrix, sheet="")

    return ParsedDocument(
        rows=rows,
        meta={"encoding": encoding, "row_count": len(rows)},
        warnings=warnings,
    )


# ---------------------------------------------------------------------------
# 叙述类
# ---------------------------------------------------------------------------


def _parse_pdf(path: Path, filename: str) -> ParsedDocument:
    from pypdf import PdfReader

    try:
        reader = PdfReader(path)
    except Exception as exc:  # noqa: BLE001
        raise IngestError(f"PDF 解析失败:{exc}", filename=filename) from exc

    pages: list[str] = []
    empty_pages: list[int] = []
    for i, page in enumerate(reader.pages):
        try:
            content = (page.extract_text() or "").strip()
        except Exception as exc:  # noqa: BLE001 - 单页坏掉不该毁掉整篇
            logger.warning("PDF 单页解析失败", extra=log_extra(filename=filename, page=i + 1,
                                                              error=str(exc)))
            content = ""
        if content:
            pages.append(content)
        else:
            empty_pages.append(i + 1)

    warnings: list[str] = []
    if empty_pages and len(empty_pages) == len(reader.pages):
        # 整篇都取不出文字 = 扫描件。这是**必须**给出明确提示的情况:
        # 否则用户看到的是「上传成功但一篇文档 0 个分块」,完全不知道
        # 该去做什么。pypdf 不做 OCR,这是能力边界,要如实说。
        raise IngestError(
            "这份 PDF 提取不出任何文字,很可能是扫描件或图片版。"
            "pypdf 只能读文本型 PDF,不包含 OCR。请改用文本版 PDF,"
            "或先用 OCR 工具转换后再上传。",
            filename=filename,
            pages=len(reader.pages),
        )
    if empty_pages:
        warnings.append(
            f"第 {'、'.join(map(str, empty_pages[:10]))}"
            f"{' 等' if len(empty_pages) > 10 else ''} 页没有提取到文字,"
            f"可能是图片页,这部分内容不会被检索到。"
        )

    return ParsedDocument(
        # 页与页之间用空行连接,让分块器在段落层级切开时尽量不跨页。
        text="\n\n".join(pages),
        meta={"pages": len(reader.pages), "pages_with_text": len(pages)},
        warnings=warnings,
    )


def _parse_docx(path: Path, filename: str) -> ParsedDocument:
    import docx

    try:
        document = docx.Document(path)
    except Exception as exc:  # noqa: BLE001
        raise IngestError(f"Word 解析失败:{exc}", filename=filename) from exc

    blocks: list[str] = []

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        # 保留标题层级:转成 Markdown 的 # 记号,后面就能直接复用
        # split_markdown() 的结构感知分块与面包屑逻辑,不必为 docx 再写一套。
        style = (para.style.name or "").lower()
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            blocks.append(f"{'#' * min(int(level), 6)} {text}")
        else:
            blocks.append(text)

    # 表格转成「行 → 句子」由分块层处理,但 docx 的表格也要读进来,
    # 商品说明书里的参数表往往就在这里面,漏掉等于丢了最有价值的部分。
    table_rows: list[dict[str, Any]] = []
    for t_i, table in enumerate(document.tables):
        matrix = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        table_rows.extend(_rows_from_matrix(matrix, sheet=f"表{t_i + 1}"))

    for row in table_rows:
        sentence = ", ".join(f"{k} {v}" for k, v in row.items() if v and k != "_sheet")
        if sentence:
            blocks.append(sentence + "。")

    if not blocks:
        raise IngestError("Word 文档里没有可提取的文字内容", filename=filename)

    return ParsedDocument(
        text="\n\n".join(blocks),
        meta={"paragraphs": len(document.paragraphs), "tables": len(document.tables)},
    )


def _parse_html(path: Path, filename: str) -> ParsedDocument:
    from bs4 import BeautifulSoup

    text, encoding = _decode(path.read_bytes(), filename)
    soup = BeautifulSoup(text, "lxml")

    # 必须先干掉这些标签。不删的话,页面里的 JS 和 CSS 会被当成正文
    # 一起入库——既污染向量,又会让用户在引用卡片里看到一坨代码。
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "aside"]):
        tag.decompose()

    title = (soup.title.string or "").strip() if soup.title else ""
    body = soup.get_text(separator="\n")
    # 压缩连续空行:HTML 的 get_text 会产生大量空行,直接送去分块
    # 会让很多 chunk 只含空白字符。
    lines = [ln.strip() for ln in body.splitlines()]
    cleaned = "\n".join(ln for ln in lines if ln)

    return ParsedDocument(
        text=cleaned,
        meta={"encoding": encoding, "title": title},
        warnings=[f"HTML 编码为 {encoding}"] if encoding == "latin-1" else [],
    )


# ---------------------------------------------------------------------------
# 入口
# ---------------------------------------------------------------------------

_PARSERS = {
    ".xlsx": _parse_xlsx,
    ".xlsm": _parse_xlsx,
    ".csv": _parse_csv,
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".html": _parse_html,
    ".htm": _parse_html,
}


def parse_document(path: Path, *, filename: str) -> ParsedDocument:
    """按扩展名分派。`filename` 是**原始文件名**,只用于选解析器和报错。"""
    kind = classify(filename)
    ext = Path(filename).suffix.lower()

    parser = _PARSERS.get(ext)
    if parser is not None:
        return parser(path, filename)

    # .md / .txt 走纯文本分支
    if kind == "text":
        return _read_text_file(path, filename)

    raise IngestError(f"没有为 {ext} 注册解析器", filename=filename)  # pragma: no cover
